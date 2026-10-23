"""Audit source seal counts, independent alpha images and editable footer text."""
import json
import argparse
from pathlib import Path
import sys
import cv2
import numpy as np
import pymupdf
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'editable_engine'))
import scan_tables
import editable_converter as engine

parser=argparse.ArgumentParser()
parser.add_argument('source',type=Path)
parser.add_argument('docx',type=Path)
parser.add_argument('without_stamps',nargs='?')
parser.add_argument('--expected-counts',default='1,2')
parser.add_argument('--expected-date',default='2026年9月17日')
args=parser.parse_args()
source_path,docx_path = args.source,args.docx
source_counts,source_pixels = [],0
with pymupdf.open(source_path) as source:
    for page in source:
        pix = page.get_pixmap(dpi=300)
        rgb = np.frombuffer(pix.samples,np.uint8).reshape(pix.height,pix.width,pix.n)
        stamps = scan_tables.extract_stamps(rgb)
        source_counts.append(len(stamps))
        for stamp in stamps:
            im = cv2.imdecode(np.frombuffer(stamp['png'],np.uint8),cv2.IMREAD_UNCHANGED)
            source_pixels += int(np.count_nonzero(im[:,:,3]))

doc = Document(docx_path)
counts = [0]*len(doc.sections)
pixels,ids = 0,[]
for block,index,_ in engine.body_blocks_by_section(doc):
    for drawing in block.iter(qn('w:drawing')):
        properties = next(drawing.iter(qn('wp:docPr')))
        ids.append(properties.get('id'))
        seal = properties.get('name','').startswith('PDF stamp ')
        blip = next(drawing.iter(qn('a:blip')))
        part = doc.part.related_parts[blip.get(qn('r:embed'))]
        im = cv2.imdecode(np.frombuffer(part.blob,np.uint8),cv2.IMREAD_UNCHANGED)
        if seal:
            counts[index] += 1
            assert im.shape[2] == 4, 'A seal needs its own transparent PNG'
            visible = im[:,:,3] > 0
            color = im[:,:,:3].astype(np.int16)
            red = (color[:,:,2]>120)&(color[:,:,2]-color[:,:,1]>35)&(color[:,:,2]-color[:,:,0]>25)
            assert np.all(red[visible]), 'Printed black text is baked into the seal'
            assert np.any(~visible), 'Seal background must be transparent'
            pixels += int(np.count_nonzero(visible))
        else:
            color = im[:,:,:3].astype(np.int16)
            red = (color[:,:,2]>120)&(color[:,:,2]-color[:,:,1]>35)&(color[:,:,2]-color[:,:,0]>25)
            assert np.count_nonzero(red) < 10, 'Seal duplicated in a background band'
assert counts == source_counts == [int(n) for n in args.expected_counts.split(',')], (counts,source_counts)
assert pixels == source_pixels, (pixels,source_pixels)
assert len(ids) == len(set(ids)), 'Duplicate Word drawing IDs'
text = '\n'.join(p.text for p in doc.paragraphs)
assert '天津鑫地机械设备有限公司' in text
assert args.expected_date in text
if args.without_stamps:
    for drawing in list(doc.element.iter(qn('w:drawing'))):
        props = next(drawing.iter(qn('wp:docPr')))
        if props.get('name','').startswith('PDF stamp '):
            drawing.getparent().remove(drawing)
    assert args.expected_date in '\n'.join(p.text for p in doc.paragraphs)
    doc.save(args.without_stamps)
print(json.dumps({'stamp_objects_by_page':counts,'duplicate_stamps':0,
                  'source_stamp_pixels':source_pixels,'output_stamp_pixels':pixels,
                  'editable_company_and_date':True,'unique_drawing_ids':True}))
