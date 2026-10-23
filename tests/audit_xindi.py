"""Independent transcription of numeric columns in the three source scans."""
from decimal import Decimal
import json
import sys
from docx import Document

doc=Document(sys.argv[1])
assert [(len(t.rows),len(t.columns)) for t in doc.tables]==[(33,7),(41,7),(39,7)]
rows=list(doc.tables[0].rows)[1:]+list(doc.tables[1].rows)[1:]+list(doc.tables[2].rows)[1:37]
quantities=[1,10,2,2,2,2,4,1,2,2,2,4,1,1,1,1,4,2,4,4,2,2,2,2,4,2,2,4,1,1,1,1,
            2,4,1,1,1,6,3,3,1,6,1,1,1,1,6,1,1,2,2,2,2,1,1,2,1,1,24,9,80,16,40,8,4,32,4,12,30,8,12,16,
            10,86,6,8,2,80,8,8,24,28,4,2,4,8,10,100,2,20,2,6,36,40,24,2,64,48,10,8,4,24,8,8,8,4,16,16]
prices=('8600 50 140 382 200 460 35 1500 350 90 28 28 1530 85 285 16 310 820 30 103 100 30 160 42 24 24 105 35 240 50 40 248 '
        '94 80 33 278 260 164 70 20 95 140 1800 1940 285 40 16 3500 540 300 40 45 34 40 40 3400 3800 2800 0.2 0.5 0.3 0.1 0.5 1.14 1 0.4 0.3 2 2 0.89 0.5 2 '
        '25.6 0.2 15.2 1.75 0.15 0.46 0.35 0.6 0.24 0.75 0.4 5 4 0.5 3 0.07 0.575 3.2 14.9 1 0.2 0.22 0.3 0.6 0.8 0.14 3 0.065 2 0.53 0.02 0.1 2 4 0.15 0.15').split()
# Amounts transcribed independently, not calculated from the tested output.
amounts=('8600 500 280 764 400 920 140 1500 700 180 56 112 1530 85 285 16 1240 1640 120 412 200 60 320 84 96 48 210 140 240 50 40 248 '
         '188 320 33 278 260 984 210 60 95 840 1800 1940 285 40 96 3500 540 600 80 90 68 40 40 6800 3800 2800 4.80 4.50 24 1.60 20 9.12 4 12.80 1.20 24 60 7.12 6 32 '
         '256 17.20 91.20 14 0.30 36.80 2.80 4.80 5.76 21 1.60 10 16 4 30 7 1.15 64 29.80 6 7.20 8.80 7.20 1.20 51.20 6.72 30 0.52 8 12.72 0.16 0.80 16 16 2.40 2.40').split()
assert len(rows)==len(quantities)==len(prices)==len(amounts)==108
for i,row in enumerate(rows):
    cells=[c.text.strip() for c in row.cells]
    assert cells[0]==str(i+1),(i+1,'sequence',cells)
    for col,value in [(3,str(quantities[i])),(4,prices[i]),(5,amounts[i])]:
        assert Decimal(cells[col])==Decimal(value),(i+1,col,cells[col],value)
    if 59<=i+1<=105:assert not cells[2],(i+1,'optional drawing should stay blank')
assert rows[92].cells[1].text.replace(' ','')=='标准型弹性垫圈5×1.3'
assert rows[100].cells[1].text.replace(' ','')=='圆螺母M25*1.5'
assert doc.tables[2].cell(37,5).text=='48004.87'
assert doc.tables[2].cell(38,1).text=='肆万捌仟零肆元捌角柒分'
assert doc.tables[2].cell(38,1)._tc is doc.tables[2].cell(38,6)._tc
assert len(doc.sections)==3
print(json.dumps({'data_rows':108,'numeric_cells_checked':432,'source_total':'48004.87','audit':'passed'},ensure_ascii=False))
