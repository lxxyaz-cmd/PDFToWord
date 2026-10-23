"""Geometry regressions for editable image tables (no OCR service required)."""
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'editable_engine'))
import cv2
import numpy as np
from docx import Document
import editable_converter as engine
import scan_tables


def fixture():
    image = np.full((700, 1000, 3), 255, np.uint8)
    xs, ys = [60, 200, 650, 930], [130, 230, 330, 430, 530]
    for y in ys:
        cv2.line(image, (xs[0], y), (xs[-1], y), (0, 0, 0), 2)
    for x in xs:
        cv2.line(image, (x, ys[0]), (x, ys[-1]), (0, 0, 0), 2)
    # Last row has a horizontal merged cell across columns 2 and 3.
    cv2.line(image, (650, 434), (650, 526), (255, 255, 255), 6)
    labels = [['No', 'Item', 'Amount'], ['1', 'Part-A', '120.00'],
              ['2', 'Part-B', '35.50'], ['Total', '155.50', '']]
    words = []
    for row, values in enumerate(labels):
        for col, text in enumerate(values):
            if not text:
                continue
            x, y = xs[col] + 20, ys[row] + 40
            cv2.putText(image, text, (x, y+20), cv2.FONT_HERSHEY_SIMPLEX, .65, (0, 0, 0), 1)
            words.append({'t': text, 'x': x, 'y': y, 'w': len(text)*13, 'h': 22})
    # OCR is allowed to return the entire page as a single line; cell assignment
    # must use word geometry, not WinRT's line.text or iteration order.
    return image, {'width': 1000, 'height': 700, 'lines': [{'words': list(reversed(words))}]}


class ScanTableTests(unittest.TestCase):
    def test_curved_photo_grid(self):
        from photo_tables import rectify_photo_table
        image=np.full((750,900,3),240,np.uint8)
        xx=np.arange(70,801)
        for row in range(16):
            yy=130+row*32+15*np.sin((xx-70)/160)
            cv2.polylines(image,[np.c_[xx,yy].round().astype(np.int32)],False,(20,20,20),2)
        for x in np.linspace(70,800,6):
            offset=15*np.sin((x-70)/160)
            cv2.line(image,(round(x),round(130+offset)),(round(x),round(610+offset)),(20,20,20),2)
        diagnostics={}
        flat=rectify_photo_table(image,diagnostics)
        self.assertIsNotNone(flat)
        self.assertEqual((diagnostics['rows'],diagnostics['columns']),(15,5))

    def test_clipped_ring_and_detached_star_are_one_seal(self):
        image = np.full((700,1000,3),255,np.uint8)
        cv2.ellipse(image,(1000,350),(90,90),0,90,270,(220,20,20),3)
        cv2.circle(image,(985,350),10,(220,20,20),-1)
        stamps = scan_tables.extract_stamps(image)
        self.assertEqual(len(stamps),1)
        transparent = cv2.imdecode(np.frombuffer(stamps[0]['png'],np.uint8),cv2.IMREAD_UNCHANGED)
        source_red = (image[:,:,0]-image[:,:,1].astype(np.int16))>35
        self.assertEqual(np.count_nonzero(transparent[:,:,3]),np.count_nonzero(source_red))

    def test_grid_words_and_merge(self):
        image, payload = fixture()
        grids = scan_tables.detect_tables(image)
        self.assertEqual(len(grids), 1)
        self.assertEqual((len(grids[0]['ys'])-1, len(grids[0]['xs'])-1), (4, 3))
        doc = Document()
        count, assigned = scan_tables.rebuild_page(doc, image, payload, 500, 700,
            engine.repair_ocr_text, engine.set_table_geometry, engine.set_run_font)
        self.assertEqual((count, assigned), (1, 11))
        table = doc.tables[0]
        self.assertEqual(table.cell(1, 1).text, 'Part-A')
        self.assertEqual(table.cell(1, 2).text, '120.00')
        self.assertEqual(table.cell(2, 2).text, '35.50')
        self.assertIs(table.cell(3, 1)._tc, table.cell(3, 2)._tc)
        self.assertEqual(table.cell(3, 1).text, '155.50')

    def test_plain_text_is_not_a_grid(self):
        image = np.full((700, 1000, 3), 255, np.uint8)
        for row in range(8):
            cv2.putText(image, 'Some ordinary text 120.00', (70, 100+row*60),
                        cv2.FONT_HERSHEY_SIMPLEX, 1, (0,0,0), 2)
        self.assertEqual(scan_tables.detect_tables(image), [])

    def test_geometry_mismatch_preserves_original(self):
        image, payload = fixture()
        payload['width'] = 500
        doc = Document()
        self.assertEqual(scan_tables.rebuild_page(doc, image, payload, 500, 700,
            engine.repair_ocr_text, engine.set_table_geometry, engine.set_run_font), (0, 0))
        self.assertEqual(len(doc.tables), 0)

    def test_decimal_repair(self):
        self.assertEqual(engine.repair_ocr_text('465．8'), '465.8')
        self.assertEqual(engine.repair_ocr_text('200，0'), '200.0')
        self.assertEqual(engine.repair_ocr_text('1，000'), '1，000')


if __name__ == '__main__':
    unittest.main()
