"""Run only the final single EXE in an isolated directory, then audit OOXML."""
import json
import hashlib
from datetime import datetime, timezone
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
from docx import Document

root = Path(__file__).resolve().parents[1]
isolated = root / 'tmp' / 'final-exe-run'
isolated.mkdir(parents=True, exist_ok=True)
executable = isolated / 'PDFToWord-Single.exe'
shutil.copy2(root / 'dist' / executable.name, executable)
environment = os.environ.copy()
environment.pop('PYTHONPATH', None)
cases = [('quotation', root/'报价单.pdf', root/'output'/'报价单-可编辑表格.docx', 'auto'),
         ('xindi',root/'天津鑫地报价单.pdf',root/'output'/'天津鑫地报价单-修复.docx','auto'),
         ('photo',root/'1.pdf',root/'output'/'1-修复.docx','auto'),
         ('mixed', root/'tmp'/'pdfs'/'mixed-image-table.pdf', root/'output'/'mixed-image-table.docx', 'editable'),
         ('native', root/'tmp'/'pdfs'/'pua-bank-table.pdf', root/'output'/'pua-bank-table.docx', 'auto')]
report_path=root/'output'/'final-exe-audit.json'
def sha256(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream,'sha256').hexdigest()
report = {'_run':{'started_at_utc':datetime.now(timezone.utc).isoformat(),'status':'running',
                  'executable_sha256':sha256(executable)}}
def save_report():
    report_path.write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
def failed_run(kind,error,traceback):
    report['_run'].update(status='failed',error=str(error))
    save_report()
    sys.__excepthook__(kind,error,traceback)
sys.excepthook=failed_run
save_report()
for label, source, destination, mode in cases:
    started=time.perf_counter()
    # A stale output must not masquerade as a successful executable run.
    test_output = isolated / (label + '.docx')
    if test_output.exists():
        test_output.unlink()
    result = subprocess.run([str(executable), '--convert', str(source), str(test_output), '--engine', mode],
                            cwd=isolated, env=environment, capture_output=True, timeout=300)
    assert result.returncode == 0 and test_output.exists(), (label, result.returncode, result.stderr)
    doc = Document(test_output)
    if label == 'quotation':
        subprocess.run([sys.executable, str(root/'tests'/'audit_quotation.py'), str(test_output)], check=True)
        subprocess.run([sys.executable, str(root/'tests'/'audit_stamps.py'), str(source),str(test_output)],check=True)
    elif label == 'photo':
        subprocess.run([sys.executable,str(root/'tests'/'audit_photo_one.py'),str(test_output)],check=True)
    elif label == 'xindi':
        subprocess.run([sys.executable,str(root/'tests'/'audit_xindi.py'),str(test_output)],check=True)
        subprocess.run([sys.executable,str(root/'tests'/'audit_stamps.py'),str(source),str(test_output),
                        '--expected-counts','1,1,2','--expected-date','2026年1月26日'],check=True)
    elif label == 'mixed':
        assert len(doc.tables) == 1 and len(doc.tables[0].columns) == 7
        assert '12345' in '\n'.join(p.text for p in doc.paragraphs)
        assert '2026' in '\n'.join(p.text for p in doc.paragraphs)
    elif label == 'native':
        text = '\n'.join(node.text or '' for node in doc.element.iter() if node.tag.endswith('}t'))
        assert doc.tables and all(value in text for value in ['100', '500', '1500'])
        assert not any(0xE000 <= ord(char) <= 0xF8FF for char in text)
    shutil.copy2(test_output, destination)
    report[label] = {'tables':len(doc.tables), 'sections':len(doc.sections), 'bytes':test_output.stat().st_size,
                     'mode':mode, 'exit_code':result.returncode,'elapsed_seconds':round(time.perf_counter()-started,2),
                     'source_sha256':sha256(source),'docx_sha256':sha256(test_output)}
    save_report()
report['_run'].update(status='passed',finished_at_utc=datetime.now(timezone.utc).isoformat())
save_report()
print(json.dumps(report,ensure_ascii=False))
