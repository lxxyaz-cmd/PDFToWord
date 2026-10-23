"""Independent transcription audit for the retained photographed 1.pdf."""
import json
from pathlib import Path
import sys
from docx import Document
from docx.oxml.ns import qn

doc=Document(sys.argv[1])
assert len(doc.tables)==1
table=doc.tables[0]
assert (len(table.rows),len(table.columns))==(42,5)
rows=table.rows[1:]
quantities=[2,4,2,4,2,2,2,2,2,2,2,2,4,4,4,4,1,2,2,1,4,2,2,1,20,20,8,6,2,15,8,10,2,4,8,8,8,4,4,4,1]
amounts=[2360,1080,860,240,180,30,220,52,360,3240,558,1120,1552,360,80,240,380,9000,240,150,45,410,90,705,9600,1600,3290,1890,192,1680,7345,415,519,1250,10260,2240,2800,880,940,132,108]
codes=[]
for section,values in [('12',range(1,5)),('11',range(8,12)),('11',range(1,8)),('12',[5]),
                       ('13',range(1,6)),('14',range(1,3)),('15',[5]),('20',range(1,12)),('15',range(1,5))]:
    codes.extend(f'JZSY1003-{section}-00-{value:03}' for value in values)
codes.extend(['GN617.1-M20×1.5','JZSY1003-90-00-007'])
assert len(codes)==len(quantities)==len(amounts)==41
for index,row in enumerate(rows):
    cells=[cell.text.strip() for cell in row.cells]
    assert cells[0]==str(index+1),(index+1,'sequence',cells[0])
    assert cells[1],(index+1,'missing name')
    assert cells[2]==codes[index],(index+1,'drawing',cells[2],codes[index])
    assert cells[3]==str(quantities[index]),(index+1,'quantity',cells[3],quantities[index])
    assert cells[4]==str(amounts[index]),(index+1,'amount',cells[4],amounts[index])
assert [(s.page_width.pt,s.page_height.pt) for s in doc.sections]==[(595.,842.)]
stamps=[node for node in doc.element.iter(qn('wp:docPr')) if node.get('name','').startswith('PDF stamp ')]
assert len(stamps)==1,'Preserve the grey side stamp as an independent picture'
print(json.dumps({'file':str(Path(sys.argv[1])),'tables':1,'rows':41,'columns':5,'drawing_numbers_checked':41,
                  'quantities_checked':41,'amounts_checked':41,'stamp_objects':1,'audit':'passed'},ensure_ascii=False))
