"""Audit the retained two-page quotation; never infer OCR values from totals."""
from decimal import Decimal
import json
from pathlib import Path
import sys
from docx import Document

path = Path(sys.argv[1])
doc = Document(path)
assert len(doc.tables) == 2, 'Expected two editable tables'
assert [(len(t.rows), len(t.columns)) for t in doc.tables] == [(27, 7), (27, 7)]
rows = list(doc.tables[0].rows)[1:] + list(doc.tables[1].rows)[1:25]
assert [r.cells[0].text.strip() for r in rows] == [str(i) for i in range(1, 51)]
# Visually transcribed from the retained source scan, not from the OCR output.
prices = ['19.6','200','30','120','3300','240','240','260','120','130','230','120','180',
          '960','350','350','240','320','310','60','15','160','330','90','150','10',
          '180','60','160','35','160','860','670','90','200','190','110','25','190',
          '350','175','465.8','360','80','400','55','180','1400','450','190']
codes = ['CZCT1200-12-00-001','CZCT6200-10-20-004','CZCT6200-10-20-005','CZCT6200A-10-70-001',
         'CZCT7001-01-10-000']
codes += ['CZCT7001-01-20-' + suffix for suffix in ['001','002','003','004','005','006','007','008','009','110','210','310','410','510']]
codes += ['CZCT7001-01-30-' + suffix for suffix in ['001','101','111','201','202','203','204','205','301','302']]
codes += ['CZCT7001-01-40-001','CZCT7001-01-40-002','CZCT7001-01-40-101']
codes += ['CZCT7001-02-00-' + suffix for suffix in ['001','002','003','004']]
codes += ['CZCT7001-02-10-' + suffix for suffix in ['101','102','103','104','106','107','108','109','110','111','112','201']]
codes += ['CZCT7001-02-30-001','CZCT7001-02-30-002']
quantities = [1]*50
for index in [1,4,20,21,22,25,26,27,28,35,36,39]:
    quantities[index-1] = 2
quantities[37] = 4
for i, row in enumerate(rows):
    cells = [cell.text.strip() for cell in row.cells]
    assert cells[2] == codes[i], (i+1, 'drawing', cells[2], codes[i])
    assert Decimal(cells[3]) == quantities[i], (i+1, 'quantity', cells[3])
    assert Decimal(cells[4]) == Decimal(prices[i]), (i+1, 'price', cells[4], prices[i])
    assert Decimal(cells[5]) == Decimal(prices[i])*quantities[i], (i+1, 'amount', cells[5])
assert doc.tables[1].cell(25,5).text == '17000.00'
assert doc.tables[1].cell(26,1).text == '壹万柒仟元整'
assert doc.tables[1].cell(26,1)._tc is doc.tables[1].cell(26,6)._tc
print(json.dumps({'tables':2, 'sections':len(doc.sections), 'data_rows':50, 'drawing_numbers_checked':50,
                  'quantities_checked':50, 'unit_prices_checked':50, 'line_amounts_checked':50,
                  'source_total':'17000.00','editable_table_audit':'passed'}, ensure_ascii=False))
