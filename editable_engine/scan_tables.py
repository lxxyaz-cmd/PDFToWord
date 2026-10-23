"""Recover ruled raster tables using local image geometry and Windows OCR.

Coordinates stay in rendered-image pixels until the Word layout is created.
No OCR service or model download is used at conversion time.
"""
from __future__ import annotations

from bisect import bisect_right
import io
from statistics import median

import cv2
import numpy as np
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt


def clusters(values, gap=8):
    groups = []
    for value in sorted(values):
        if not groups or value - groups[-1][-1] > gap:
            groups.append([value])
        else:
            groups[-1].append(value)
    return [int(round(sum(group) / len(group))) for group in groups]


def extract_stamps(image):
    """One transparent object per spatial stamp, including disconnected stars.

    Group on the whole page before table/band slicing. The grouping mask is
    deliberately wider than the export mask: no dilated pixels are exported.
    """
    red, green, blue = [image[:, :, i].astype(np.int16) for i in range(3)]
    mask = ((red > 120) & (red-green > 35) & (red-blue > 25)).astype(np.uint8)
    kernel = max(15, round(min(image.shape[:2]) * .027)) | 1
    blobs = cv2.dilate(mask, np.ones((kernel, kernel), np.uint8))
    contours, _ = cv2.findContours(blobs, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    stamps = []
    boxes = [cv2.boundingRect(contour) for contour in contours]
    for index,(x,y,w,h) in enumerate(boxes):
        # An edge-clipped circular seal has an open ring. Its detached star
        # may still be a separate contour; include it in the surrounding ring's
        # crop exactly once instead of exporting a second floating star.
        if any(j != index and ow*oh > w*h and ox <= x and oy <= y
               and ox+ow >= x+w and oy+oh >= y+h
               for j,(ox,oy,ow,oh) in enumerate(boxes)):
            continue
        yy, xx = np.nonzero(mask[y:y+h, x:x+w])
        if len(xx) < 100:
            continue
        x0, y0, x1, y1 = x+int(xx.min()), y+int(yy.min()), x+int(xx.max())+1, y+int(yy.max())+1
        if x1-x0 < 30 or y1-y0 < 30:
            continue
        rgba = cv2.cvtColor(image[y0:y1, x0:x1], cv2.COLOR_RGB2BGRA)
        rgba[:, :, 3] = mask[y0:y1, x0:x1] * 255
        stamps.append({'bbox': (x0,y0,x1,y1), 'png': cv2.imencode('.png', rgba)[1].tobytes()})
    return sorted(stamps, key=lambda stamp: (stamp['bbox'][1],stamp['bbox'][0]))


def stamp_free_image(image, stamps):
    """Recover black scan text using the red channel; leave stamp ink separate."""
    clean = image.copy()
    for stamp in stamps:
        x0,y0,x1,y1 = stamp['bbox']
        # Black printed letters remain dark in the red channel. Red stamp ink
        # becomes white, including its anti-aliased edges, without inpainting or
        # inventing characters that are obscured in the source.
        gray = np.minimum(image[y0:y1,x0:x1,0].astype(np.float32) * (255/120),255).astype(np.uint8)
        clean[y0:y1,x0:x1] = gray[:, :, None]
    return clean


def add_stamp(paragraph, stamp, scale, x_offset, y_offset, number):
    x0,y0,x1,y1 = stamp['bbox']
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(stamp['png']), width=Pt((x1-x0)*scale))
    drawing = run._r.find(qn('w:drawing'))
    inline = drawing.find(qn('wp:inline'))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "r")} distT="0" distB="0" distL="0" distR="0" '
        'simplePos="0" relativeHeight="10" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:posOffset>{int(x_offset*scale*12700)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{int(y_offset*scale*12700)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int((x1-x0)*scale*12700)}" cy="{int((y1-y0)*scale*12700)}"/>'
        '<wp:wrapNone/></wp:anchor>')
    for tag in ('wp:docPr','wp:cNvGraphicFramePr','a:graphic'):
        element = inline.find(qn(tag))
        if element is not None:
            if tag == 'wp:docPr':
                element.set('name',f'PDF stamp {number}')
                element.set('descr','Independent transparent stamp extracted from source PDF')
            anchor.append(element)
    drawing.remove(inline)
    drawing.append(anchor)


def detect_tables(image):
    """Find connected line grids; text strokes alone must not become a table."""
    gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    ink = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                cv2.THRESH_BINARY_INV, 31, 15)
    height, width = gray.shape
    horizontal = cv2.morphologyEx(ink, cv2.MORPH_OPEN,
                                cv2.getStructuringElement(cv2.MORPH_RECT, (max(24, width // 35), 1)))
    vertical = cv2.morphologyEx(ink, cv2.MORPH_OPEN,
                              cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(40, width // 30, height // 45))))
    # Close small scan/JPEG gaps without joining separate nearby tables.
    horizontal = cv2.morphologyEx(horizontal, cv2.MORPH_CLOSE, np.ones((1, 9), np.uint8))
    vertical = cv2.morphologyEx(vertical, cv2.MORPH_CLOSE, np.ones((9, 1), np.uint8))
    grid = cv2.bitwise_or(horizontal, vertical)
    joined = cv2.dilate(grid, np.ones((5, 5), np.uint8))
    contours, _ = cv2.findContours(joined, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    tables = []
    for contour in contours:
        left, top, box_width, box_height = cv2.boundingRect(contour)
        if box_width < width * .15 or box_height < 45:
            continue
        right, bottom = min(width, left + box_width), min(height, top + box_height)
        h = horizontal[top:bottom, left:right]
        v = vertical[top:bottom, left:right]
        # Individual lines may end at merged cells, so count their total support.
        xs = clusters((np.where(np.count_nonzero(v, axis=0) > max(25, box_height * .18))[0] + left).tolist())
        ys = clusters((np.where(np.count_nonzero(h, axis=1) > max(25, box_width * .25))[0] + top).tolist())
        if len(xs) < 3 or len(ys) < 3 or len(xs) > 65 or len(ys) > 300:
            continue
        if min(np.diff(xs)) < 15 or min(np.diff(ys)) < 12:
            continue
        tables.append({"xs": xs, "ys": ys, "horizontal": horizontal,
                       "vertical": vertical, "bbox": (xs[0], ys[0], xs[-1], ys[-1])})
    return sorted(tables, key=lambda item: (item["ys"][0], item["xs"][0]))


def word_box(word):
    x, y = float(word.get("x", 0)), float(word.get("y", 0))
    return x, y, x + float(word.get("w", 0)), y + float(word.get("h", 0))


def words_text(words, repair):
    """Rebuild reading lines within ONE cell, even if OCR joined several columns."""
    if not words:
        return ""
    line_height = max(6, median(float(word.get("h", 1)) for word in words) * .55)
    lines = []
    for word in sorted(words, key=lambda w: (float(w["y"]), float(w["x"]))):
        center = float(word["y"]) + float(word["h"]) / 2
        target = next((line for line in lines if abs(line[0] - center) < line_height), None)
        if target is None:
            target = [center, []]
            lines.append(target)
        target[1].append(word)
    result = []
    for _, line in sorted(lines):
        ordered = sorted(line, key=lambda w: float(w["x"]))
        # A wide whitespace gap is meaningful for prose but not Chinese glyphs
        # or punctuation within a number / drawing identifier.
        text, previous = "", None
        for word in ordered:
            value = str(word.get("t", ""))
            if previous is not None:
                gap = float(word["x"]) - (float(previous["x"]) + float(previous["w"]))
                if gap > max(3, float(word["h"]) * .3):
                    text += " "
            text += value
            previous = word
        result.append(repair(text))
    return "\n".join(result)


def cell_groups(table):
    """Merge only rectangular groups where a ruled internal border is absent."""
    xs, ys = table["xs"], table["ys"]
    rows, cols = len(ys) - 1, len(xs) - 1
    parent = list(range(rows * cols))

    def root(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def join(a, b):
        parent[root(b)] = root(a)

    def coverage(mask, x0, y0, x1, y1, axis):
        crop = mask[max(0, y0):y1, max(0, x0):x1]
        return float(np.mean(np.any(crop > 0, axis=axis))) if crop.size else 1.

    for row in range(rows):
        for col in range(cols):
            i = row * cols + col
            if col + 1 < cols and coverage(table["vertical"], xs[col + 1] - 5, ys[row] + 6,
                                          xs[col + 1] + 6, ys[row + 1] - 5, 1) < .25:
                join(i, i + 1)
            if row + 1 < rows and coverage(table["horizontal"], xs[col] + 6, ys[row + 1] - 5,
                                          xs[col + 1] - 5, ys[row + 1] + 6, 0) < .25:
                join(i, i + cols)
    grouped = {}
    for i in range(rows * cols):
        grouped.setdefault(root(i), []).append(divmod(i, cols))
    result = []
    for members in grouped.values():
        r0, c0 = min(r for r, _ in members), min(c for _, c in members)
        r1, c1 = max(r for r, _ in members), max(c for _, c in members)
        if len(members) == (r1 - r0 + 1) * (c1 - c0 + 1):
            result.append((r0, c0, r1, c1))
        else:
            result.extend((r, c, r, c) for r, c in members)
    return result


def add_word_table(document, table, words, scale, repair, set_geometry, set_font):
    xs, ys = table["xs"], table["ys"]
    rows, cols = len(ys) - 1, len(xs) - 1
    result = document.add_table(rows=rows, cols=cols)
    result.style = "Table Grid"
    set_geometry(result, [(b - a) * scale for a, b in zip(xs, xs[1:])])
    groups = cell_groups(table)
    cell_words = {(r, c): [] for r, c, _, _ in groups}
    owners = {}
    for r0, c0, r1, c1 in groups:
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                owners[r, c] = (r0, c0)
    assigned = set()
    for index, word in enumerate(words):
        x0, y0, x1, y1 = word_box(word)
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        if xs[0] <= cx <= xs[-1] and ys[0] <= cy <= ys[-1]:
            row = min(rows - 1, bisect_right(ys, cy) - 1)
            col = min(cols - 1, bisect_right(xs, cx) - 1)
            cell_words[owners[row, col]].append(word)
            assigned.add(index)
    font_size = max(7., min(11., median([float(w["h"]) for w in words] or [30]) * scale * .9))
    for row_index, row in enumerate(result.rows):
        row.height = Pt((ys[row_index + 1] - ys[row_index]) * scale * .9)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    for r0, c0, r1, c1 in groups:
        cell = result.cell(r0, c0)
        if r1 != r0 or c1 != c0:
            cell = cell.merge(result.cell(r1, c1))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins = OxmlElement('w:tcMar')
        properties = cell._tc.get_or_add_tcPr()
        old_margins = properties.find(qn('w:tcMar'))
        if old_margins is not None:
            properties.remove(old_margins)
        for side, value in [('top',24),('bottom',24),('start',60),('end',60)]:
            node = OxmlElement('w:'+side)
            node.set(qn('w:w'), str(value))
            node.set(qn('w:type'), 'dxa')
            margins.append(node)
        properties.append(margins)
        cell.text = words_text(cell_words[r0, c0], repair)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.
            # The paragraph mark otherwise inherits the 11 pt Normal style,
            # forcing every 7-9 pt CJK row taller and spilling scan pages.
            ppr = paragraph._p.get_or_add_pPr()
            rpr = OxmlElement('w:rPr')
            for tag in ('w:sz', 'w:szCs'):
                size = OxmlElement(tag)
                size.set(qn('w:val'), str(round(font_size*2)))
                rpr.append(size)
            ppr.append(rpr)
            for run in paragraph.runs:
                set_font(run)
                run.font.size = Pt(font_size)
    table['word_table'] = result
    return assigned


def rebuild_page(document, image, payload, usable_width, usable_height,
                 repair, set_geometry, set_font):
    """Append a fully editable fragment, or leave the original scan untouched."""
    tables = detect_tables(image)
    words = [word for line in payload.get("lines", []) for word in line.get("words", [])
             if str(word.get("t", "")).strip()]
    if not tables or not words:
        return 0, 0
    # Coordinates must correspond to the exact image passed to Windows OCR.
    height, width = image.shape[:2]
    if payload.get("width") != width or payload.get("height") != height:
        return 0, 0
    scale = min(usable_width / width, usable_height / height)
    stamps = extract_stamps(image)
    clean_image = stamp_free_image(image, stamps)
    stamps.extend(payload.get('photo_side_artwork',[]))
    consumed = set()
    elements = []
    for table in tables:
        x0, y0, x1, y1 = table["bbox"]
        inside = [word for word in words if x0 <= (word_box(word)[0] + word_box(word)[2]) / 2 <= x1
                  and y0 <= (word_box(word)[1] + word_box(word)[3]) / 2 <= y1]
        # A photo / drawing grid without recognised cell text is not a table.
        if len(inside) < 2:
            continue
        elements.append((y0, "table", table))
    if not elements:
        return 0, 0
    table_boxes = [entry[2]["bbox"] for entry in elements]
    # Keep non-table artwork (letterheads, signatures, seals) as image bands.
    # Those pixels are not table data and must not be replaced by uncertain OCR.
    previous = 0
    for y0, y1 in sorted((box[1], box[3]) for box in table_boxes):
        if y0 > previous:
            elements.append((previous, "image", (previous, y0)))
        previous = max(previous, y1)
    if previous < height:
        elements.append((previous, "image", (previous, height)))
    band_holders = []
    for _, kind, data in sorted(elements, key=lambda entry: entry[0]):
        if kind == "table":
            consumed.update(add_word_table(document, data, words, scale, repair, set_geometry, set_font))
        elif kind == "image":
            y0, y1 = data
            band_stamps = [stamp for stamp in stamps if y0 <= (stamp['bbox'][1]+stamp['bbox'][3])/2 < y1]
            if band_stamps:
                # Printed company/date text must be real Word text, not baked
                # into the seal picture. OCR ran on the stamp-free source.
                outside = [word for word in words if y0 <= float(word['y'])+float(word['h'])/2 < y1]
                lines = []
                for word in sorted(outside,key=lambda w: float(w['y'])):
                    cy = float(word['y'])+float(word['h'])/2
                    line = next((line for line in lines if abs(line[0]-cy) < max(8,float(word['h'])*.6)),None)
                    if line is None:
                        line = [cy,[]]
                        lines.append(line)
                    line[1].append(word)
                holder, reference_y = None, y0
                for cy,line in sorted(lines):
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_before = Pt(max(6,(min(float(w['y']) for w in line)-y0)*scale) if holder is None else 0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.
                    paragraph.paragraph_format.left_indent = Pt(min(usable_width*.8,min(float(w['x']) for w in line)*scale))
                    run = paragraph.add_run(words_text(line,repair))
                    set_font(run)
                    run.font.size = Pt(max(8,min(11,median(float(w['h']) for w in line)*scale)))
                    if holder is None:
                        holder, reference_y = paragraph,min(float(w['y']) for w in line)
                if holder is None:
                    holder = document.add_paragraph()
                band_holders.append((y0,y1,holder,reference_y))
                # Keep space for a floating seal even when no source text was
                # recognised, so it cannot overlap the next section.
                reserve = max(stamp['bbox'][3] for stamp in band_stamps)-reference_y
                last = document.paragraphs[-1]
                last.paragraph_format.space_after = Pt(max(6,reserve*scale-18))
                continue
            crop = clean_image[y0 + 5:max(y0 + 5, y1 - 5), 12:-12]
            if not crop.size:
                continue
            ink = np.min(crop, axis=2) < 170
            # Ignore isolated dust and the scanner's dark paper edge.
            rows = np.where(np.count_nonzero(ink, axis=1) > max(8, width * .008))[0]
            if not len(rows):
                continue
            crop = crop[max(0, rows[0]-5):min(crop.shape[0], rows[-1]+6)]
            encoded = cv2.imencode(".png", cv2.cvtColor(crop, cv2.COLOR_RGB2BGR))[1].tobytes()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.
            paragraph.add_run().add_picture(io.BytesIO(encoded), width=Pt(crop.shape[1] * scale))
    for number,stamp in enumerate(stamps,1):
        x0,y0,x1,y1 = stamp['bbox']
        cy = (y0+y1)/2
        owner = next((entry[2] for entry in elements if entry[1]=='table' and entry[2]['ys'][0] <= cy <= entry[2]['ys'][-1]),None)
        if owner is not None:
            xs,ys = owner['xs'],owner['ys']
            row = min(len(ys)-2,max(0,bisect_right(ys,y0)-1))
            col = min(len(xs)-2,max(0,bisect_right(xs,x0)-1))
            paragraph = owner['word_table'].cell(row,col).paragraphs[0]
            add_stamp(paragraph,stamp,scale,x0-xs[col],y0-ys[row],number)
        else:
            holder = next((band for band in band_holders if band[0] <= cy < band[1]),None)
            if holder is not None:
                before = holder[2].paragraph_format.space_before
                # Word/Writer anchor at the paragraph box, before its leading
                # spacing. Include that spacing to keep the seal below the
                # table and aligned with the independently editable text.
                leading_pixels = (before.pt if before is not None else 0)/scale
                add_stamp(holder[2],stamp,scale,x0,y0-holder[3]+leading_pixels,number)
    return len(table_boxes), len(consumed)
